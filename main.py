import sqlite3
from datetime import datetime, timedelta, time
from openpyxl import Workbook
import csv
import os

DB = "restaurant.db"
OPEN_TIME = time(11, 0)
CLOSE_TIME = time(22, 0)

# ---------- DB ----------
def db():
    return sqlite3.connect(DB)

def setup():
    with db() as con:
        con.execute("""
        CREATE TABLE IF NOT EXISTS tables(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            seats INTEGER NOT NULL
        )""")
        con.execute("""
        CREATE TABLE IF NOT EXISTS reservations(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            guests INTEGER NOT NULL,
            table_id INTEGER NOT NULL,
            start DATETIME NOT NULL,
            end DATETIME NOT NULL
        )""")

# ---------- TABLES ----------
def add_table():
    seats = int(input("Seats: "))
    with db() as con:
        con.execute("INSERT INTO tables(seats) VALUES(?)", (seats,))
    print("Table added.")

def list_tables():
    with db() as con:
        rows = con.execute("SELECT id, seats FROM tables ORDER BY seats").fetchall()
    for r in rows:
        print(f"Table {r[0]} – {r[1]} seats")

# ---------- RESERVATIONS ----------
def is_open(start, end):
    return OPEN_TIME <= start.time() and end.time() <= CLOSE_TIME

def available(table_id, start, end):
    with db() as con:
        cur = con.execute("""
        SELECT 1 FROM reservations
        WHERE table_id = ?
        AND NOT (end <= ? OR start >= ?)
        """, (table_id, start, end))
        return cur.fetchone() is None

def find_table(guests, start, end):
    with db() as con:
        tables = con.execute(
            "SELECT id FROM tables WHERE seats >= ? ORDER BY seats", (guests,)
        ).fetchall()
    for (tid,) in tables:
        if available(tid, start, end):
            return tid
    return None

def reserve():
    name = input("Name: ").strip()
    guests = int(input("Guests: "))
    start = datetime.fromisoformat(input("Start (YYYY-MM-DD HH:MM): "))
    duration = int(input("Duration (minutes): "))
    end = start + timedelta(minutes=duration)

    if not is_open(start, end):
        print("Outside opening hours.")
        return

    table_id = find_table(guests, start, end)
    if not table_id:
        print("No table available.")
        return

    with db() as con:
        con.execute("""
        INSERT INTO reservations(name, guests, table_id, start, end)
        VALUES(?,?,?,?,?)
        """, (name, guests, table_id, start, end))

    print(f"Reserved table {table_id} for {name}.")

def list_reservations(date_filter=None):
    q = """
    SELECT id, name, guests, table_id, start, end
    FROM reservations
    """
    params = ()
    if date_filter:
        q += " WHERE date(start)=?"
        params = (date_filter,)
    q += " ORDER BY start"

    with db() as con:
        rows = con.execute(q, params).fetchall()

    if not rows:
        print("No reservations.")
        return

    for r in rows:
        print(f"[{r[0]}] {r[1]} | {r[2]} guests | Table {r[3]} | {r[4]} → {r[5]}")

def search_by_name():
    name = input("Search name: ").strip()
    with db() as con:
        rows = con.execute("""
        SELECT id, name, start, table_id
        FROM reservations
        WHERE name LIKE ?
        """, (f"%{name}%",)).fetchall()
    for r in rows:
        print(r)

def delete_reservation():
    rid = input("Reservation ID: ")
    with db() as con:
        con.execute("DELETE FROM reservations WHERE id=?", (rid,))
    print("Deleted.")

# ---------- EXPORT ----------
def excel_it():
    with db() as con:
        rows = con.execute("SELECT * FROM reservations ORDER BY start").fetchall()

    if not rows:
        print("Nothing to export.")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "Reservations"
    ws.append(["ID", "Name", "Guests", "Table", "Start", "End"])
    for r in rows:
        ws.append(r)

    path = os.path.join(os.getcwd(), "reservations.xlsx")
    wb.save(path)
    print(f"Excel exported: {path}")

def csv_export():
    with db() as con:
        rows = con.execute("SELECT * FROM reservations").fetchall()

    with open("reservations.csv", "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["ID", "Name", "Guests", "Table", "Start", "End"])
        writer.writerows(rows)

    print("CSV exported.")

# ---------- CLI ----------
def menu():
    print("\n--- Restaurant ---")
    print("1) Add table")
    print("2) List tables")
    print("3) New reservation")
    print("4) List reservations")
    print("5) List reservations by date")
    print("6) Search by name")
    print("7) Delete reservation")
    print("8) Export Excel")
    print("9) Export CSV")
    print("0) Exit")

def main():
    setup()
    while True:
        menu()
        c = input("> ")
        if c == "1":
            add_table()
        elif c == "2":
            list_tables()
        elif c == "3":
            reserve()
        elif c == "4":
            list_reservations()
        elif c == "5":
            list_reservations(input("Date (YYYY-MM-DD): "))
        elif c == "6":
            search_by_name()
        elif c == "7":
            delete_reservation()
        elif c == "8":
            excel_it()
        elif c == "9":
            csv_export()
        elif c == "0":
            break

if __name__ == "__main__":
    main()

